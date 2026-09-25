# 独立实现的公共物理模型（Claude 审查复算用）。
# 只读原始 Excel/DEM 副本；不导入冻结项目的 src 代码。
# 口径：论文式(2)-(15)、(17)-(21)、(35)-(38)；双精度浮点。
import math, os
import numpy as np
import openpyxl, tifffile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
RAW = os.path.join(ROOT, 'inputs', 'raw')
G0 = 9.80665


def _rows(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    return [r for r in wb.worksheets[0].iter_rows(values_only=True)]


def load_data():
    base = os.path.join(RAW, 'base')
    rows = _rows(os.path.join(base, '调度中心与服务区.xlsx'))
    nodes = {}
    for r in rows:
        if r[0] and (str(r[0]).startswith('O0') or str(r[0]).startswith('S0')) and isinstance(r[2], (int, float)):
            nodes[r[0]] = dict(lon=float(r[2]), lat=float(r[3]), alt=float(r[4]))
    wb = openpyxl.load_workbook(os.path.join(base, '物资需求与配送时限.xlsx'), data_only=True)
    ws = wb['逐箱货箱清单']
    boxes = {}
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if not r[0]:
            continue
        first = (r[5] == '是')
        boxes[r[0]] = dict(site=r[1], cat=r[2], mass=float(r[3]), vol=float(r[4]), first=first,
                           first_dl=(float(r[6]) if r[6] is not None else None), expect=float(r[7]), w=float(r[8]))
    for b in boxes.values():
        hard = []
        if b['cat'] == '医疗物资':
            hard.append(b['expect'])
        if b['first']:
            hard.append(b['first_dl'])
        b['hard'] = min(hard) if hard else None
    rows = _rows(os.path.join(base, '运输无人机数据.xlsx'))
    types = {}
    for r in rows:
        if r[0] in ('A', 'B', 'C') and isinstance(r[1], str) and isinstance(r[2], (int, float)):
            types[r[0]] = dict(m0=float(r[2]), Q=float(r[3]), V=float(r[4]), vc=float(r[5]), L0=float(r[6]), LF=float(r[7]),
                               E=float(r[8]), rho=float(r[9]) / 100, prep=float(r[10]), load=float(r[11]), hb=float(r[12]),
                               hbox=float(r[13]), vup=float(r[14]), vdn=float(r[15]), eta=float(r[16]))
    for r in rows:
        if r[0] in ('A', 'B', 'C') and isinstance(r[1], (int, float)) and r[2] is not None and not isinstance(r[3], (int, float)):
            types[r[0]]['nbat'] = int(r[1]); types[r[0]]['Tfull'] = float(r[2])
    drones = {r[0]: r[1] for r in rows if r[0] and str(r[0]).startswith('U0')}
    rows = _rows(os.path.join(base, '中继无人机数据.xlsx'))
    relay = None
    for r in rows:
        if r[0] == 'R' and isinstance(r[1], str) and isinstance(r[2], (int, float)):
            relay = dict(m=float(r[4]), vc=float(r[5]), Pc=float(r[6]), E=float(r[7]), rho=float(r[8]) / 100, prep=float(r[9]),
                         link=float(r[10]), turn=float(r[11]), vup=float(r[12]), vdn=float(r[13]), eta=float(r[14]),
                         Ph=float(r[16]), Pcom=float(r[17]), agl_max=float(r[18]))
        elif r[0] == 'R' and isinstance(r[1], (int, float)):
            relay['ncomp'] = int(r[1]); relay['Tfull'] = float(r[2])
    rows = _rows(os.path.join(base, '通信链路参数.xlsx'))
    link = {}
    for r in rows[2:]:
        if r[0] is None:
            continue
        link[(r[0], r[1])] = float(r[4])
    return nodes, boxes, types, drones, relay, link


class DEM:
    def __init__(self):
        t = tifffile.TiffFile(os.path.join(RAW, 'dem.tif'))
        p = t.pages[0]
        self.h = p.asarray().astype(np.float64)
        sx, sy, _ = p.tags['ModelPixelScaleTag'].value
        tp = p.tags['ModelTiepointTag'].value
        gk = p.tags['GeoKeyDirectoryTag'].value
        # GTRasterTypeGeoKey(1025)=2 -> PixelIsPoint：tiepoint 为像元(0,0)中心
        keys = {gk[4 + 4 * i]: gk[4 + 4 * i + 3] for i in range(gk[3])}
        assert keys.get(1025) == 2, keys
        self.dx, self.dy = sx, sy
        self.lon_edge0 = tp[3] - sx / 2
        self.lat_edge0 = tp[4] + sy / 2
        self.nr, self.nc = self.h.shape

    def uv(self, lon, lat):
        return (np.asarray(lon) - self.lon_edge0) / self.dx, (self.lat_edge0 - np.asarray(lat)) / self.dy

    def cell_value(self, lon, lat):
        u, v = self.uv(lon, lat)
        j, i = int(math.floor(u)), int(math.floor(v))
        return self.h[i, j]

    def segment_cells(self, lon0, lat0, lon1, lat1, tol=1e-9):
        # 与线段相交（含只擦过边界/角点）的全部像元闭方格
        u0, v0 = self.uv(lon0, lat0); u1, v1 = self.uv(lon1, lat1)
        ts = {0.0, 1.0}
        for a0, a1 in ((u0, u1), (v0, v1)):
            if abs(a1 - a0) > 0:
                lo, hi = sorted((a0, a1))
                for k in range(math.ceil(lo), math.floor(hi) + 1):
                    ts.add((k - a0) / (a1 - a0))
        ts = sorted(t for t in ts if 0 <= t <= 1)
        cells = set()
        def add_point(u, v):
            for jj in {math.floor(u + tol), math.floor(u - tol)}:
                for ii in {math.floor(v + tol), math.floor(v - tol)}:
                    cells.add((ii, jj))
        for t in ts:
            add_point(u0 + t * (u1 - u0), v0 + t * (v1 - v0))
        for a, b in zip(ts[:-1], ts[1:]):
            tm = (a + b) / 2
            add_point(u0 + tm * (u1 - u0), v0 + tm * (v1 - v0))
        for (i, j) in cells:
            assert 0 <= i < self.nr and 0 <= j < self.nc, ('outside DEM', i, j)
        return cells


class Proj:
    def __init__(self, lon0, lat0):
        a = 6378137.0; f = 1 / 298.257223563; e2 = f * (2 - f)
        s = math.sin(math.radians(lat0))
        N = a / math.sqrt(1 - e2 * s * s); M = a * (1 - e2) / (1 - e2 * s * s) ** 1.5
        self.kx = math.pi / 180 * N * math.cos(math.radians(lat0)); self.ky = math.pi / 180 * M
        self.lon0, self.lat0 = lon0, lat0

    def xy(self, lon, lat):
        return (np.asarray(lon) - self.lon0) * self.kx, (np.asarray(lat) - self.lat0) * self.ky

    def dist3(self, p, q):
        x0, y0 = self.xy(p[0], p[1]); x1, y1 = self.xy(q[0], q[1])
        return math.sqrt((x1 - x0) ** 2 + (y1 - y0) ** 2 + (q[2] - p[2]) ** 2)


class Model:
    def __init__(self):
        self.nodes, self.boxes, self.types, self.drones, self.relay, self.link = load_data()
        self.dem = DEM()
        o = self.nodes['O01']
        self.proj = Proj(o['lon'], o['lat'])
        self._geo = {}

    def op_alt(self, node):
        n = self.nodes[node]
        return n['alt'] if node.startswith('O') else n['alt'] + 30.0

    def geometry(self, a_ll, b_ll, za, zb):
        key = (tuple(a_ll), tuple(b_ll), za, zb)
        if key in self._geo:
            return self._geo[key]
        cells = self.dem.segment_cells(a_ll[0], a_ll[1], b_ll[0], b_ll[1])
        tmax = max(self.dem.h[i, j] for i, j in cells)
        zc = tmax + 50.0
        x0, y0 = self.proj.xy(*a_ll); x1, y1 = self.proj.xy(*b_ll)
        d = math.hypot(float(x1 - x0), float(y1 - y0))
        g = dict(distance=d, terrain_max=tmax, cruise_z=zc, up=zc - za, down=zc - zb, cells=len(cells))
        self._geo[key] = g
        return g

    def leg_geo(self, i, j):
        a, b = self.nodes[i], self.nodes[j]
        return self.geometry((a['lon'], a['lat']), (b['lon'], b['lat']), self.op_alt(i), self.op_alt(j))

    def leg(self, t, i, j, q):
        g = self.types[t]; geo = self.leg_geo(i, j)
        L = g['L0'] - (g['L0'] - g['LF']) * (q / g['Q']) ** 1.5
        eh = g['E'] * geo['distance'] / L
        eu = (g['m0'] + q) * G0 * geo['up'] / g['eta'] / 3.6e6
        tm = geo['up'] / g['vup'] + geo['distance'] / g['vc'] + geo['down'] / g['vdn']
        return dict(time=tm, energy=eh + eu, eh=eh, eu=eu, L=L, geo=geo)

    @staticmethod
    def tchg(s, T):
        return T * (0.65 * (0.9 - s) / 0.9 + 0.35) if s < 0.9 else T * 0.35 * (1 - s) / 0.1

    def route(self, t, boxes, order, start=0.0):
        """按给定访问顺序重算架次：事件、交付完成、能耗、SOC、返航时刻"""
        g = self.types[t]
        by = {}
        for b in boxes:
            by.setdefault(self.boxes[b]['site'], []).append(b)
        assert set(by) == set(order), (by.keys(), order)
        q = sum(self.boxes[b]['mass'] for b in boxes)
        vol = sum(self.boxes[b]['vol'] for b in boxes)
        clock = start + g['prep'] + g['load'] * len(boxes)
        takeoff = clock
        events = [('prep', start, clock, None, None)]
        E = 0.0; deliv = {}; prev = 'O01'
        path = list(order) + ['O01']
        for nxt in path:
            lg = self.leg(t, prev, nxt, q)
            geo = lg['geo']
            a, b = self.nodes[prev], self.nodes[nxt]
            za, zb, zc = self.op_alt(prev), self.op_alt(nxt), geo['cruise_z']
            t1 = clock + geo['up'] / g['vup']; t2 = t1 + geo['distance'] / g['vc']; t3 = t2 + geo['down'] / g['vdn']
            events.append(('up', clock, t1, (a['lon'], a['lat'], za), (a['lon'], a['lat'], zc)))
            events.append(('cruise', t1, t2, (a['lon'], a['lat'], zc), (b['lon'], b['lat'], zc)))
            events.append(('down', t2, t3, (b['lon'], b['lat'], zc), (b['lon'], b['lat'], zb)))
            E += lg['energy']; clock = t3
            if nxt != 'O01':
                n = len(by[nxt]); th = clock + g['hb'] + g['hbox'] * n
                events.append(('handover', clock, th, (b['lon'], b['lat'], zb), (b['lon'], b['lat'], zb)))
                for bx in by[nxt]:
                    deliv[bx] = th
                q -= sum(self.boxes[bx]['mass'] for bx in by[nxt]); clock = th
            prev = nxt
        soc = 1 - E / g['E']
        return dict(type=t, mass=sum(self.boxes[b]['mass'] for b in boxes), vol=vol, energy=E, soc=soc, ret=clock,
                    takeoff=takeoff, deliveries=deliv, events=events, start=start,
                    bat_free=clock + self.tchg(soc, g['Tfull']))

    # ---- 中继 ----
    def relay_mission(self, pos, start, service_end):
        r = self.relay; o = self.nodes['O01']
        geo = self.geometry((o['lon'], o['lat']), (pos[0], pos[1]), o['alt'], pos[2])
        zc = geo['cruise_z']
        up_out = zc - o['alt']; down_out = zc - pos[2]
        up_back = zc - pos[2]; down_back = zc - o['alt']
        t_out = up_out / r['vup'] + geo['distance'] / r['vc'] + down_out / r['vdn']
        t_back = up_back / r['vup'] + geo['distance'] / r['vc'] + down_back / r['vdn']
        arrival = start + r['prep'] + t_out
        sstart = arrival + r['link']
        ret = service_end + t_back
        e_cruise = r['Pc'] * (2 * geo['distance'] / r['vc']) / 3600
        e_up = r['m'] * G0 * (up_out + up_back) / r['eta'] / 3.6e6
        e_hover = (r['Ph'] + r['Pcom']) * (service_end - arrival) / 3600
        E = e_cruise + e_up + e_hover
        soc = 1 - E / r['E']
        agl = pos[2] - self.dem.cell_value(pos[0], pos[1])
        return dict(arrival=arrival, service_start=sstart, ret=ret, energy=E, soc=soc, down_out=down_out,
                    entity_free=ret + r['turn'], comp_free=ret + r['turn'] + self.tchg(soc, r['Tfull']), agl=agl,
                    geo=geo, t_out=t_out, t_back=t_back)

    # ---- 链路 ----
    def lmax(self, kind):
        L = self.link
        sens = L[('接收参数', '接收灵敏度（dBm）')] + L[('接收参数', '衰落裕量（dB）')]
        sys_ = L[('传播参数', '系统损耗（dB）')]
        T = (L[('运输无人机', '发射功率（dBm）')], L[('运输无人机', '天线增益（dBi）')])
        A = (L[('中继接入端', '发射功率（dBm）')], L[('中继接入端', '天线增益（dBi）')])
        B = (L[('中继回传端', '发射功率（dBm）')], L[('中继回传端', '天线增益（dBi）')])
        G = (L[('固定网关 G01', '发射功率（dBm）')], L[('固定网关 G01', '天线增益（dBi）')])
        pairs = {'TG': (T, G), 'TA': (T, A), 'BG': (B, G)}
        x, y = pairs[kind]
        f = lambda a, b: a[0] + a[1] + b[1] - sys_ - sens
        return min(f(x, y), f(y, x))

    def fspl(self, dkm):
        f = self.link[('传播参数', '载波频率（MHz）')]
        return 32.45 + 20 * math.log10(f) + 20 * math.log10(dkm)

    def gateway(self):
        o = self.nodes['O01']
        return (o['lon'], o['lat'], o['alt'] + self.link[('固定网关 G01', '天线离地高度（m）')])
